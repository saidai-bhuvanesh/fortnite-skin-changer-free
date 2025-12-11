"""
Document Processor - Extract text from PDF, DOCX, TXT files
"""

import os
from pathlib import Path
from typing import List, Dict, Optional
from loguru import logger
import sys

# Import document processing libraries
try:
    import PyPDF2
    import pdfplumber
except ImportError:
    logger.warning("PDF libraries not installed")

try:
    from docx import Document
except ImportError:
    logger.warning("python-docx not installed")

# Configure logger
logger.remove()
logger.add(sys.stdout, level="INFO")
logger.add("logs/document_processor.log", rotation="10 MB")

class DocumentProcessor:
    """Process various document formats for RAG ingestion"""
    
    def __init__(self):
        self.supported_formats = [".pdf", ".docx", ".txt", ".md"]
    
    def process_document(self, file_path: str) -> Dict[str, any]:
        """
        Process a single document and extract text
        
        Args:
            file_path: Path to document file
            
        Returns:
            Dict with extracted text and metadata
        """
        file_path = Path(file_path)
        
        if not file_path.exists():
            logger.error(f"File not found: {file_path}")
            return {"text": "", "metadata": {}, "error": "File not found"}
        
        file_ext = file_path.suffix.lower()
        
        if file_ext not in self.supported_formats:
            logger.warning(f"Unsupported format: {file_ext}")
            return {"text": "", "metadata": {}, "error": f"Unsupported format: {file_ext}"}
        
        try:
            if file_ext == ".pdf":
                text = self._extract_pdf(file_path)
            elif file_ext == ".docx":
                text = self._extract_docx(file_path)
            elif file_ext in [".txt", ".md"]:
                text = self._extract_text(file_path)
            else:
                text = ""
            
            metadata = {
                "filename": file_path.name,
                "file_type": file_ext,
                "file_size": file_path.stat().st_size,
                "path": str(file_path)
            }
            
            logger.info(f"Processed {file_path.name}: {len(text)} characters")
            
            return {
                "text": text,
                "metadata": metadata,
                "error": None
            }
        
        except Exception as e:
            logger.error(f"Error processing {file_path}: {e}")
            return {
                "text": "",
                "metadata": {"filename": file_path.name},
                "error": str(e)
            }
    
    def _extract_pdf(self, file_path: Path) -> str:
        """Extract text from PDF using pdfplumber (better for tables/layout)"""
        text_parts = []
        
        try:
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(page_text)
        except Exception as e:
            logger.warning(f"pdfplumber failed, trying PyPDF2: {e}")
            # Fallback to PyPDF2
            try:
                with open(file_path, 'rb') as file:
                    reader = PyPDF2.PdfReader(file)
                    for page in reader.pages:
                        text_parts.append(page.extract_text())
            except Exception as e2:
                logger.error(f"Both PDF extractors failed: {e2}")
                raise
        
        return "\n\n".join(text_parts)
    
    def _extract_docx(self, file_path: Path) -> str:
        """Extract text from DOCX file"""
        doc = Document(file_path)
        text_parts = []
        
        # Extract paragraphs
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)
        
        # Extract tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text for cell in row.cells)
                text_parts.append(row_text)
        
        return "\n".join(text_parts)
    
    def _extract_text(self, file_path: Path) -> str:
        """Extract text from plain text files"""
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
    
    def process_directory(self, directory_path: str) -> List[Dict[str, any]]:
        """
        Process all supported documents in a directory
        
        Args:
            directory_path: Path to directory containing documents
            
        Returns:
            List of processed document dicts
        """
        directory = Path(directory_path)
        
        if not directory.exists():
            logger.error(f"Directory not found: {directory}")
            return []
        
        documents = []
        
        for file_path in directory.rglob("*"):
            if file_path.is_file() and file_path.suffix.lower() in self.supported_formats:
                doc = self.process_document(str(file_path))
                if doc["text"]:  # Only add if text was extracted
                    documents.append(doc)
        
        logger.info(f"Processed {len(documents)} documents from {directory}")
        return documents
    
    def chunk_text(
        self,
        text: str,
        chunk_size: int = 1000,
        chunk_overlap: int = 200
    ) -> List[str]:
        """
        Split text into overlapping chunks for better RAG retrieval
        
        Args:
            text: Input text to chunk
            chunk_size: Target chunk size in characters
            chunk_overlap: Overlap between chunks
            
        Returns:
            List of text chunks
        """
        if len(text) <= chunk_size:
            return [text]
        
        chunks = []
        start = 0
        
        while start < len(text):
            end = start + chunk_size
            
            # Try to break at sentence boundary
            if end < len(text):
                # Look for sentence ending near chunk boundary
                for delimiter in ['. ', '.\n', '! ', '?\n', '? ']:
                    last_delimiter = text.rfind(delimiter, start, end)
                    if last_delimiter != -1:
                        end = last_delimiter + 1
                        break
            
            chunk = text[start:end].strip()
            if chunk:
                chunks.append(chunk)
            
            # Move start position with overlap
            start = end - chunk_overlap
        
        logger.debug(f"Split text into {len(chunks)} chunks")
        return chunks
